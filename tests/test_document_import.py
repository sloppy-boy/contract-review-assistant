from io import BytesIO
import subprocess
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
import pytest

from app.document_import import LocalTesseractOCR, extract_document


def test_extract_txt_preserves_paragraph_locations_and_source_metadata():
    source = "合同编号：HT-001\n甲方：示例甲公司\n乙方：示例乙公司\n签订日期：2026-09-07\n合同金额：人民币100元\n\n第一条 付款。"
    result = extract_document("notice.txt", source.encode("utf-8"))

    assert result["text"] == source
    assert result["segments"] == [
        {"id": "paragraph-1", "location": {"paragraph": 1}, "text": "合同编号：HT-001\n甲方：示例甲公司\n乙方：示例乙公司\n签订日期：2026-09-07\n合同金额：人民币100元"},
        {"id": "paragraph-2", "location": {"paragraph": 2}, "text": "第一条 付款。"},
    ]
    assert result["metadata"]["filename"] == {"value": "notice.txt", "source": "upload.filename"}
    for key, value, quote in [
        ("contractNumber", "HT-001", "合同编号：HT-001"),
        ("partyA", "示例甲公司", "甲方：示例甲公司"),
        ("partyB", "示例乙公司", "乙方：示例乙公司"),
        ("signedDate", "2026-09-07", "签订日期：2026-09-07"),
        ("amount", "人民币100元", "合同金额：人民币100元"),
    ]:
        assert result["metadata"][key] == {
            "value": value,
            "source": {
                "type": "segment",
                "segmentId": "paragraph-1",
                "location": {"paragraph": 1},
                "quote": quote,
            },
        }
    assert result["warnings"] == []
    assert result["method"] == "text"


def _docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("主合同")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "服务费"
    table.cell(1, 1).text = "100元"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_extract_docx_locates_paragraphs_and_table_rows():
    result = extract_document("contract.DOCX", _docx_bytes())

    assert result["method"] == "docx"
    assert result["text"] == "主合同\n项目 | 金额\n服务费 | 100元"
    assert result["segments"] == [
        {"id": "paragraph-1", "location": {"paragraph": 1}, "text": "主合同"},
        {
            "id": "table-1-row-1",
            "location": {"table": 1, "row": 1},
            "text": "项目 | 金额",
        },
        {
            "id": "table-1-row-2",
            "location": {"table": 1, "row": 2},
            "text": "服务费 | 100元",
        },
    ]


def _text_pdf(*page_texts: str) -> bytes:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{' '.join(f'{4 + i * 2} 0 R' for i in range(len(page_texts)))}] /Count {len(page_texts)} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for index, page_text in enumerate(page_texts):
        content_id = 5 + index * 2
        content = f"BT /F1 12 Tf 72 720 Td ({page_text}) Tj ET".encode("ascii")
        objects.extend(
            [
                f"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 3 0 R >> >> /MediaBox [0 0 612 792] /Contents {content_id} 0 R >>".encode(),
                b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
            ]
        )
    result = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for object_id, body in enumerate(objects, 1):
        offsets.append(len(result))
        result.extend(f"{object_id} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = len(result)
    result.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    result.extend(b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:]))
    result.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(result)


def test_extract_pdf_preserves_page_locations():
    result = extract_document("agreement.pdf", _text_pdf("First page", "Second page"))

    assert result["method"] == "pdf-text"
    assert result["text"] == "First page\nSecond page"
    assert result["segments"] == [
        {"id": "page-1", "location": {"page": 1}, "text": "First page"},
        {"id": "page-2", "location": {"page": 2}, "text": "Second page"},
    ]


def test_scanned_pdf_requires_an_ocr_engine():
    from pypdf import PdfWriter

    output = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(output)

    with pytest.raises(RuntimeError, match="OCR is required"):
        extract_document("scan.pdf", output.getvalue())


def test_scanned_pdf_uses_supplied_ocr_callback_with_page_location():
    from pypdf import PdfWriter

    output = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(output)
    calls = []

    def ocr(pdf_data: bytes, page_number: int) -> str:
        calls.append((pdf_data, page_number))
        return "扫描文字"

    result = extract_document("scan.pdf", output.getvalue(), ocr=ocr)

    assert calls == [(output.getvalue(), 1)]
    assert result["method"] == "pdf-ocr"
    assert result["segments"] == [
        {"id": "page-1", "location": {"page": 1}, "text": "扫描文字"}
    ]
    assert result["warnings"] == ["page 1 text was extracted with OCR"]


@pytest.mark.parametrize(
    ("filename", "data"),
    [("empty.txt", b""), ("broken.docx", b"not-a-zip"), ("broken.pdf", b"not-a-pdf")],
)
def test_empty_and_damaged_documents_are_rejected(filename, data):
    with pytest.raises(ValueError):
        extract_document(filename, data)


def test_raw_upload_size_is_limited_to_ten_mibibytes():
    with pytest.raises(ValueError, match="10 MiB"):
        extract_document("large.txt", b"x" * (10 * 1024 * 1024 + 1))


def test_extracted_text_is_limited_to_one_million_characters():
    with pytest.raises(ValueError, match="1,000,000"):
        extract_document("verbose.txt", b"x" * 1_000_001)


def test_docx_uncompressed_content_is_limited_to_fifty_mibibytes():
    source = BytesIO(_docx_bytes())
    output = BytesIO()
    with ZipFile(source) as original, ZipFile(output, "w", ZIP_DEFLATED) as expanded:
        for item in original.infolist():
            expanded.writestr(item, original.read(item.filename))
        expanded.writestr("word/oversized.bin", b"x" * (50 * 1024 * 1024 + 1))

    with pytest.raises(ValueError, match="50 MiB"):
        extract_document("expanded.docx", output.getvalue())


def test_pdf_page_count_is_limited_to_one_hundred():
    from pypdf import PdfWriter

    output = BytesIO()
    writer = PdfWriter()
    for _ in range(101):
        writer.add_blank_page(width=10, height=10)
    writer.write(output)

    with pytest.raises(ValueError, match="100 pages"):
        extract_document("long.pdf", output.getvalue())


def test_local_tesseract_reports_a_missing_binary_as_runtime_error(tmp_path):
    engine = LocalTesseractOCR(
        rasterizer_cmd=str(tmp_path / "missing-pdftoppm"),
        tesseract_cmd=str(tmp_path / "missing-tesseract"),
        timeout=0.1,
    )

    with pytest.raises(RuntimeError, match="local OCR dependency is unavailable"):
        engine(b"synthetic PDF bytes", 1)


def test_markdown_is_extracted_as_located_plain_text():
    result = extract_document("terms.md", b"# Terms\n\nPayment is due.")

    assert result["method"] == "text"
    assert [segment["text"] for segment in result["segments"]] == ["# Terms", "Payment is due."]


def test_docx_preserves_body_order_across_paragraphs_and_tables():
    document = Document()
    document.add_paragraph("Before")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Term"
    table.cell(0, 1).text = "Value"
    document.add_paragraph("After")
    output = BytesIO()
    document.save(output)

    result = extract_document("ordered.docx", output.getvalue())

    assert [segment["text"] for segment in result["segments"]] == [
        "Before", "Term | Value", "After"
    ]
    assert [segment["id"] for segment in result["segments"]] == [
        "paragraph-1", "table-1-row-1", "paragraph-2"
    ]


def test_docx_with_malformed_document_xml_is_a_value_error():
    source, output = BytesIO(_docx_bytes()), BytesIO()
    with ZipFile(source) as original, ZipFile(output, "w", ZIP_DEFLATED) as damaged:
        for item in original.infolist():
            content = b"<w:document" if item.filename == "word/document.xml" else original.read(item.filename)
            damaged.writestr(item, content)

    with pytest.raises(ValueError, match="invalid or damaged DOCX"):
        extract_document("damaged.docx", output.getvalue())


def test_mixed_pdf_requires_ocr_for_its_textless_pages():
    with pytest.raises(RuntimeError, match="page 2.*OCR"):
        extract_document("mixed.pdf", _text_pdf("Text page", ""))


def test_mixed_pdf_rejects_an_ocr_page_that_produces_no_text():
    with pytest.raises(RuntimeError, match="page 2.*no text"):
        extract_document("mixed.pdf", _text_pdf("Text page", ""), ocr=lambda _data, _page: "")


def test_local_ocr_failure_does_not_expose_stderr(monkeypatch):
    def fail(*_args, **_kwargs):
        raise subprocess.CalledProcessError(1, ["ocr"], stderr=b"secret document content")

    monkeypatch.setattr(subprocess, "run", fail)
    engine = LocalTesseractOCR()
    with pytest.raises(RuntimeError) as error:
        engine(b"synthetic PDF bytes", 1)

    assert str(error.value) == "local OCR failed"
    assert "secret" not in str(error.value)


@pytest.mark.parametrize("timeout", [0, -1, float("nan"), float("inf")])
def test_local_ocr_timeout_must_be_finite_and_positive(timeout):
    with pytest.raises(ValueError, match="finite and positive"):
        LocalTesseractOCR(timeout=timeout)
